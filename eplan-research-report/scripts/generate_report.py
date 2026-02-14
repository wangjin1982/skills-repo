#!/usr/bin/env python3
import sys
from datetime import datetime
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
except ImportError:
    print("Error: python-docx library is required")
    print("Install with: pip install python-docx")
    sys.exit(1)

def add_paragraph(doc, text, style='Normal', bold=False, font_size=11, spacing_after=12):
    para = doc.add_paragraph(text, style=style)
    run = para.runs[0]
    run.font.size = Pt(font_size)
    run.font.name = 'Microsoft YaHei'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    run.bold = bold
    para.paragraph_format.line_spacing = 1.5
    para.paragraph_format.space_after = Pt(spacing_after)
    return para

def add_heading(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = 'Microsoft YaHei'
        run.font.color.rgb = RGBColor(0, 51, 102)
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    return heading

def add_subheading(doc, text):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.name = 'Microsoft YaHei'
    run.font.color.rgb = RGBColor(0, 51, 102)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    para.paragraph_format.line_spacing = 1.5
    para.paragraph_format.space_after = Pt(6)
    return para

def create_cover_page(doc, client_name, report_date=None):
    if report_date is None:
        report_date = datetime.now().strftime("%Y年%m月")
    title = doc.add_paragraph()
    title_run = title.add_run('Eplan Digital Consulting')
    title_run.font.size = Pt(28)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0, 51, 102)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph()
    subtitle_run = subtitle.add_run('Research Report')
    subtitle_run.font.size = Pt(20)
    subtitle_run.font.bold = True
    subtitle_run.font.color.rgb = RGBColor(0, 51, 102)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for _ in range(5):
        doc.add_paragraph()
    client_para = doc.add_paragraph()
    client_para.add_run('Client: ').font.size = Pt(14)
    client_run = client_para.add_run(client_name)
    client_run.font.size = Pt(18)
    client_run.font.bold = True
    client_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_para = doc.add_paragraph()
    date_para.add_run('Date: ').font.size = Pt(12)
    date_run = date_para.add_run(report_date)
    date_run.font.size = Pt(12)
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

def create_report_structure(doc):
    add_heading(doc, '1. Customer Current State Analysis', level=1)
    add_paragraph(doc, '[Fill in customer current state analysis]', font_size=11)
    doc.add_paragraph()
    add_heading(doc, '2. Business Pain Points', level=1)
    add_paragraph(doc, '[Fill in business pain points analysis]', font_size=11)
    doc.add_paragraph()
    add_heading(doc, '3. Data Flow Analysis', level=1)
    add_paragraph(doc, '[Fill in data flow analysis]', font_size=11)
    doc.add_paragraph()
    add_heading(doc, '4. ROI Assumptions', level=1)
    add_subheading(doc, '4.1 Investment Cost Analysis')
    add_paragraph(doc, '[List project investment costs]', font_size=11)
    add_subheading(doc, '4.2 Expected Benefits Analysis')
    add_paragraph(doc, '[List expected benefits]', font_size=11)
    add_subheading(doc, '4.3 Quantitative Metrics')
    add_paragraph(doc, '[List quantitative metrics]', font_size=11)
    add_subheading(doc, '4.4 Investment Payback Period')
    add_paragraph(doc, '[Analyze investment payback period]', font_size=11)
    doc.add_paragraph()
    add_heading(doc, '5. Implementation Roadmap', level=1)
    add_paragraph(doc, '[Fill in implementation roadmap]', font_size=11)
    doc.add_paragraph()
    add_paragraph(doc, '— End of Report —', bold=True, font_size=11)
    conclusion = doc.add_paragraph()
    conclusion_run = conclusion.add_run('This report is based on on-site research and interviews, aiming to provide professional recommendations for your Eplan digital transformation. Specific implementation plans require further refinement based on actual conditions.')
    conclusion_run.font.size = Pt(10)
    conclusion_run.font.color.rgb = RGBColor(128, 128, 128)
    conclusion_run.font.italic = True
    conclusion.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

def generate_report(client_name, output_path, **kwargs):
    doc = Document()
    section = doc.sections[0]
    section.page_height = Inches(11.69)
    section.page_width = Inches(8.27)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    report_date = kwargs.get('report_date', None)
    create_cover_page(doc, client_name, report_date)
    create_report_structure(doc)
    output_file = Path(output_path)
    output_file.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_file))
    print(f"Report generated: {output_file}")
    return str(output_file)

if __name__ == '__main__':
    if len(sys.argv) < 3:
        print("Usage: python generate_report.py <client_name> <output_path> [report_date]")
        print("Example: python generate_report.py 'ABC Corp' './reports/ABC_Corp_report.docx'")
        sys.exit(1)
    client = sys.argv[1]
    output = sys.argv[2]
    date = sys.argv[3] if len(sys.argv) > 3 else None
    generate_report(client, output, report_date=date)
